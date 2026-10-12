from __future__ import annotations

from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
import csv, io, uuid

import openpyxl
from sqlalchemy.orm import Session

from ..models import Client, Entity, EntityChangeRequest, EntityVersion, User, UserEntityAccess
from ..timeutil import utcnow_naive

ENTITY_FIELDS = [
    'legal_name','country','currency','group_currency','characterization','company_code','erp_system',
    'tax_jurisdiction','tax_rate','customs_territory','effective_from','effective_to','active'
]

REQUIRED = ['entity_code','legal_name','country','currency','characterization','company_code','erp_system']

HEADER_MAP = {
    'Entity_ID':'entity_code','Company_Code':'company_code','Legal_Name':'legal_name','Country':'country',
    'Currency':'currency','TP_Characterization':'characterization','ERP_System':'erp_system',
    'Group_Currency':'group_currency','Tax_Jurisdiction':'tax_jurisdiction','Tax_Rate':'tax_rate',
    'Customs_Territory':'customs_territory','Effective_From':'effective_from','Effective_To':'effective_to','Active':'active'
}

def entity_snapshot(entity: Entity) -> dict:
    return {
        'entity_code':entity.entity_code,'legal_name':entity.legal_name,'country':entity.country,
        'currency':entity.currency,'group_currency':entity.group_currency,'characterization':entity.characterization,
        'company_code':entity.company_code,'erp_system':entity.erp_system,'tax_jurisdiction':entity.tax_jurisdiction,
        'tax_rate':float(entity.tax_rate) if entity.tax_rate is not None else None,
        'customs_territory':entity.customs_territory,'effective_from':entity.effective_from,
        'effective_to':entity.effective_to,'active':bool(entity.active),
    }

def _iso_date(value, field):
    if value in (None,''): return None
    text=str(value).strip()[:10]
    try: date.fromisoformat(text)
    except ValueError: raise ValueError(f'{field} must use YYYY-MM-DD')
    return text

def normalize_entity_payload(payload: dict, *, existing: Entity|None=None) -> dict:
    base=entity_snapshot(existing) if existing else {}
    base.update({k:v for k,v in payload.items() if v is not None})
    data={}
    data['entity_code']=str(base.get('entity_code','')).strip().upper()
    data['legal_name']=str(base.get('legal_name','')).strip()
    data['country']=str(base.get('country','')).strip().upper()
    data['currency']=str(base.get('currency','')).strip().upper()
    data['group_currency']=str(base.get('group_currency') or 'EUR').strip().upper()
    data['characterization']=str(base.get('characterization','')).strip()
    data['company_code']=str(base.get('company_code','')).strip()
    data['erp_system']=str(base.get('erp_system','')).strip()
    data['tax_jurisdiction']=str(base.get('tax_jurisdiction') or data['country']).strip()
    data['customs_territory']=str(base.get('customs_territory') or ('EU' if data['country'] in {'DE','FR','IT','NL','PL','ES','BE','AT','PT','IE','LU','FI','SE','DK','CZ','SK','SI','HR','RO','BG','HU','GR','CY','MT','LT','LV','EE'} else '')).strip() or None
    data['effective_from']=_iso_date(base.get('effective_from') or date.today().isoformat(),'effective_from')
    data['effective_to']=_iso_date(base.get('effective_to'),'effective_to')
    active=base.get('active', True)
    if isinstance(active,str): active=active.strip().lower() not in {'false','0','no','nein','inactive'}
    data['active']=bool(active)
    raw_tax=base.get('tax_rate')
    if raw_tax in (None,''): data['tax_rate']=None
    else:
        try:
            d=Decimal(str(raw_tax).replace('%','').replace(',','.'))
            if d>1: d=d/Decimal('100')
            if d<0 or d>1: raise ValueError
            data['tax_rate']=str(d)
        except (InvalidOperation, ValueError): raise ValueError('tax_rate must be between 0 and 1 (or 0–100%)')
    for f in REQUIRED:
        if not data.get(f): raise ValueError(f'Missing required field: {f}')
    if not (2 <= len(data['entity_code']) <= 40): raise ValueError('entity_code must contain 2–40 characters')
    if len(data['country'])!=2 or not data['country'].isalpha(): raise ValueError('country must be a 2-letter ISO code')
    if len(data['currency'])!=3 or not data['currency'].isalpha(): raise ValueError('currency must be a 3-letter ISO code')
    if len(data['group_currency'])!=3 or not data['group_currency'].isalpha(): raise ValueError('group_currency must be a 3-letter ISO code')
    if data['effective_to'] and data['effective_from'] and data['effective_to']<data['effective_from']:
        raise ValueError('effective_to cannot be before effective_from')
    return data

def validate_uniqueness(db: Session, workspace_id:int, data:dict, *, exclude_entity_id:int|None=None):
    q=db.query(Entity).filter(Entity.workspace_id==workspace_id,Entity.entity_code==data['entity_code'])
    if exclude_entity_id: q=q.filter(Entity.id!=exclude_entity_id)
    if q.first(): raise ValueError(f"Entity code {data['entity_code']} already exists")
    q=db.query(Entity).filter(Entity.workspace_id==workspace_id,Entity.erp_system==data['erp_system'],Entity.company_code==data['company_code'],Entity.active==True)
    if exclude_entity_id: q=q.filter(Entity.id!=exclude_entity_id)
    dup=q.first()
    if dup: raise ValueError(f"Company code {data['company_code']} already exists in ERP {data['erp_system']} ({dup.entity_code})")

def next_entity_version(db:Session, entity_id:int)->int:
    latest=db.query(EntityVersion).filter(EntityVersion.entity_id==entity_id).order_by(EntityVersion.version_no.desc()).first()
    return (latest.version_no+1) if latest else 1

def create_change_request(db:Session, *, workspace_id:int, user:User, change_type:str, payload:dict, reason:str|None=None, batch_code:str|None=None)->EntityChangeRequest:
    change_type=change_type.upper()
    if change_type not in {'CREATE','UPDATE','DEACTIVATE','REACTIVATE'}: raise ValueError('Unsupported change type')
    code=str(payload.get('entity_code','')).strip().upper()
    entity=db.query(Entity).filter(Entity.workspace_id==workspace_id,Entity.entity_code==code).first() if code else None
    if change_type=='CREATE':
        if entity: raise ValueError(f'Entity {code} already exists')
        data=normalize_entity_payload(payload)
        validate_uniqueness(db,workspace_id,data)
    else:
        if not entity: raise ValueError(f'Entity {code} not found')
        if change_type=='UPDATE':
            if payload.get('entity_code') and str(payload['entity_code']).upper()!=entity.entity_code: raise ValueError('Entity code is immutable; create a new entity instead')
            data=normalize_entity_payload(payload,existing=entity)
            validate_uniqueness(db,workspace_id,data,exclude_entity_id=entity.id)
        else:
            data=entity_snapshot(entity);data['active']=change_type=='REACTIVATE'
    req=EntityChangeRequest(
        workspace_id=workspace_id,entity_id=entity.id if entity else None,
        request_code=f"ECR-{code}-{uuid.uuid4().hex[:8].upper()}",entity_code=data['entity_code'],
        change_type=change_type,status='DRAFT',payload=data,reason=(reason or '').strip() or None,
        batch_code=batch_code,created_by=user.email,
    )
    db.add(req);db.flush();return req

def submit_change_request(db:Session, req:EntityChangeRequest, user:User):
    if req.status!='DRAFT': raise ValueError('Only DRAFT requests can be submitted')
    if req.created_by!=user.email and user.role!='PLATFORM_ADMIN': raise PermissionError('Only the request creator or Platform Admin can submit this draft')
    req.status='SUBMITTED';req.submitted_by=user.email;req.submitted_at=utcnow_naive();db.flush();return req

def approve_change_request(db:Session, req:EntityChangeRequest, user:User):
    if req.status!='SUBMITTED': raise ValueError('Only SUBMITTED requests can be approved')
    if req.created_by==user.email: raise PermissionError('Maker-checker control: request creator cannot approve the same master-data change')
    data=normalize_entity_payload(req.payload)
    entity=db.get(Entity,req.entity_id) if req.entity_id else None
    previous=entity_snapshot(entity) if entity else None
    if req.change_type=='CREATE':
        validate_uniqueness(db,req.workspace_id,data)
        client=db.query(Client).filter(Client.workspace_id==req.workspace_id,Client.active==True).order_by(Client.id).first()
        entity=Entity(workspace_id=req.workspace_id,client_id=client.id if client else None,
            entity_code=data['entity_code'],legal_name=data['legal_name'],country=data['country'],tax_jurisdiction=data['tax_jurisdiction'],
            currency=data['currency'],group_currency=data['group_currency'],characterization=data['characterization'],company_code=data['company_code'],
            erp_system=data['erp_system'],tax_rate=Decimal(data['tax_rate']) if data['tax_rate'] is not None else None,
            customs_territory=data['customs_territory'],active=data['active'],effective_from=data['effective_from'],effective_to=data['effective_to'])
        db.add(entity);db.flush();req.entity_id=entity.id
    else:
        if not entity: raise ValueError('Entity no longer exists')
        validate_uniqueness(db,req.workspace_id,data,exclude_entity_id=entity.id)
        for f in ENTITY_FIELDS:
            if f=='tax_rate': setattr(entity,f,Decimal(data[f]) if data[f] is not None else None)
            elif f in data: setattr(entity,f,data[f])
        db.flush()
    version=EntityVersion(workspace_id=req.workspace_id,entity_id=entity.id,version_no=next_entity_version(db,entity.id),
        snapshot=entity_snapshot(entity),effective_from=entity.effective_from,effective_to=entity.effective_to,
        approved_by=user.email,approval_request_code=req.request_code)
    db.add(version)
    # Give a Local Finance creator access to an entity they onboarded. Other enterprise roles use workspace-wide policy.
    creator=db.query(User).filter(User.email==req.created_by).first()
    if creator and creator.role=='LOCAL_FINANCE' and not db.query(UserEntityAccess).filter(UserEntityAccess.user_id==creator.id,UserEntityAccess.entity_id==entity.id).first():
        db.add(UserEntityAccess(user_id=creator.id,entity_id=entity.id,can_read=True,can_write=True,can_approve=False))
    req.status='APPROVED';req.reviewed_by=user.email;req.reviewed_at=utcnow_naive();db.flush()
    return req,entity,previous,entity_snapshot(entity),version.version_no

def reject_change_request(db:Session, req:EntityChangeRequest, user:User, reason:str):
    if req.status!='SUBMITTED': raise ValueError('Only SUBMITTED requests can be rejected')
    if req.created_by==user.email: raise PermissionError('Maker-checker control: request creator cannot review the same master-data change')
    req.status='REJECTED';req.reviewed_by=user.email;req.reviewed_at=utcnow_naive();req.reason=(reason or req.reason or '').strip() or 'Rejected during review';db.flush();return req

def parse_entity_master(filename:str, content:bytes)->list[dict]:
    name=filename.lower()
    rows=[]
    if name.endswith('.xlsx'):
        wb=openpyxl.load_workbook(BytesIO(content),read_only=True,data_only=True);ws=wb[wb.sheetnames[0]]
        iterator=ws.iter_rows(values_only=True);headers=[str(x or '').strip() for x in next(iterator)]
        for values in iterator: rows.append(dict(zip(headers,values)))
        wb.close()
    elif name.endswith('.csv'):
        text=content.decode('utf-8-sig');rows=list(csv.DictReader(io.StringIO(text)))
    else: raise ValueError('Entity Master import supports .xlsx or .csv')
    mapped=[]
    for row in rows:
        if not any(v not in (None,'') for v in row.values()): continue
        data={HEADER_MAP.get(str(k).strip(),str(k).strip()):v for k,v in row.items()}
        mapped.append(data)
    return mapped

def preview_entity_master(db:Session, *, workspace_id:int, filename:str, content:bytes)->dict:
    rows=parse_entity_master(filename,content);result=[];counts={'CREATE':0,'UPDATE':0,'NO_CHANGE':0,'INVALID':0}
    for idx,row in enumerate(rows,start=2):
        try:
            code=str(row.get('entity_code','')).strip().upper();existing=db.query(Entity).filter(Entity.workspace_id==workspace_id,Entity.entity_code==code).first() if code else None
            data=normalize_entity_payload(row,existing=existing)
            if existing:
                before=entity_snapshot(existing)
                compare={k:data.get(k) for k in data}
                changes={}
                for k in compare:
                    left,right=before.get(k),compare.get(k)
                    if k=='tax_rate':
                        left=None if left in (None,'') else Decimal(str(left));right=None if right in (None,'') else Decimal(str(right))
                    if left!=right: changes[k]={'from':before.get(k),'to':compare.get(k)}
                action='UPDATE' if changes else 'NO_CHANGE'
                if action=='UPDATE': validate_uniqueness(db,workspace_id,data,exclude_entity_id=existing.id)
            else:
                validate_uniqueness(db,workspace_id,data);changes={};action='CREATE'
            result.append({'row':idx,'entity_code':data['entity_code'],'legal_name':data['legal_name'],'action':action,'changes':changes,'payload':data,'errors':[]});counts[action]+=1
        except Exception as exc:
            result.append({'row':idx,'entity_code':str(row.get('entity_code','')),'legal_name':str(row.get('legal_name','')),'action':'INVALID','changes':{},'payload':row,'errors':[str(exc)]});counts['INVALID']+=1
    return {'filename':filename,'rows':len(result),'counts':counts,'items':result,'note':'Rows missing from the file are never interpreted as deletions/deactivations.'}
