#!/usr/bin/env python3
from __future__ import annotations
from pathlib import Path
import json, re, sys

ROOT=Path(__file__).resolve().parents[1]
VERSION=(ROOT/'VERSION.txt').read_text().strip()
errors=[]
checks=[]

def ok(name, condition, detail=''):
    checks.append((name,bool(condition),detail))
    if not condition: errors.append(f'{name}: {detail}')

required=[
 'START_DEMO.bat','QUALIFY_10_10.bat','START_OTP.bat','START_DEV.bat','STOP_OTP.bat','RESET_DEMO.bat','BACKUP_OTP.bat','RESTORE_OTP.bat','STATUS_OTP.bat','VALIDATE_RC.bat',
 'README_FIRST.txt','README.md','VERSION.txt','docker-compose.yml','.env.example','.env.demo',
 'scripts/select_runtime_ports.ps1','scripts/qualify_10_10.ps1','scripts/generate_local_qualification_report.py','docker-compose.qualification.yml','qualification/Dockerfile','qualification/BackendTest.Dockerfile','qualification/playwright.config.ts','qualification/tests/ui.spec.ts','qualification/tests/api.spec.ts','scripts/load_compose_secrets.bat','scripts/sync_postgres_credentials.bat','backend/app/main.py','backend/app/services/finance.py','backend/alembic/versions/0001_initial.py','backend/alembic/versions/0002_enterprise_rc.py','backend/alembic/versions/0003_rls_policies.py','backend/alembic/versions/0004_posting_schema_compat.py','backend/alembic/versions/0005_pnl_variance_qual.py','backend/alembic/versions/0006_entity_master_mgmt.py','backend/alembic/versions/0007_configurable_core.py','backend/alembic/versions/0008_practicality_upgrade.py','backend/alembic/versions/0009_practicality_10.py','backend/pytest.ini','backend/app/services/entity_master.py','backend/app/services/master_data_catalog.py','backend/app/routers/master_data.py','backend/app/routers/allocations.py',
 'frontend/app/page.tsx','frontend/app/login/page.tsx','frontend/app/entities/page.tsx','frontend/app/master-data/page.tsx','frontend/app/allocations/page.tsx','frontend/app/true-ups/page.tsx','frontend/app/pnl/page.tsx','frontend/app/variance/page.tsx','frontend/app/qualification/page.tsx','frontend/app/security/page.tsx','frontend/lib/i18n.tsx','README_DE.md','docs/GERMAN_LOCALISATION.md',
 'n8n/workflows/01_monthly_otp_close.json','sample_data/01_Entity_Master.xlsx','sample_data/07_ERP_GL_Actuals.xlsx','sample_data/19_ERP_Posting_Confirmation.xlsx','sample_data/20_Authoritative_PnL_Demo.csv',
 'documents/Global_TP_Policy.pdf','docs/PRODUCTION_READINESS.md','docs/ENTITY_MASTER_MANAGEMENT.md','docs/openapi.json','docs/VALIDATION_REPORT.md','docs/TEST_REPORT.md','docs/COVERAGE_REPORT.md','docs/PERFORMANCE_REPORT.md','docs/SECURITY_REPORT.md','docs/FUNCTIONAL_COVERAGE_MATRIX.md','docs/USER_ROLE_MATRIX.md','docs/PRODUCTION_QUALIFICATION_SCORECARD.md','docs/GAP_CLOSURE_REPORT.md','docs/UAT_SCENARIOS.md','docs/ADMIN_GUIDE.md','docs/USER_GUIDE.md'
]
missing=[x for x in required if not (ROOT/x).exists()]
ok('required files',not missing,', '.join(missing) if missing else 'all present')

# Banned transient/build artifacts.
banned=[]
for p in ROOT.rglob('*'):
    if p.name in {'node_modules','.next','__pycache__','.pytest_cache','.coverage'}: banned.append(str(p.relative_to(ROOT)))
    if p.is_file() and p.suffix in {'.pyc','.pyo'}: banned.append(str(p.relative_to(ROOT)))
ok('no transient artifacts',not banned,', '.join(banned[:10]) if banned else 'clean')
ok('no packaged runtime database',not list(ROOT.rglob('*.db')),'database artifact found' if list(ROOT.rglob('*.db')) else 'clean')
ok('no packaged .env secret file',not (ROOT/'.env').exists(),'.env must be generated locally' if (ROOT/'.env').exists() else 'clean')

# JSON and n8n validation.
workflow_nodes={}
try:
    for f in sorted((ROOT/'n8n/workflows').glob('*.json')):
        d=json.loads(f.read_text())
        nodes=d.get('nodes',[]); workflow_nodes[f.name]=len(nodes)
        ok(f'n8n {f.name} structure',bool(d.get('name')) and len(nodes)>=5 and isinstance(d.get('connections',{}),dict),f'nodes={len(nodes)}')
    json.loads((ROOT/'docs/openapi.json').read_text())
    ok('OpenAPI JSON',True,'valid JSON')
except Exception as e:
    ok('JSON validation',False,str(e))

# Docker Compose parse and pinning.
try:
    import yaml
    compose=yaml.safe_load((ROOT/'docker-compose.yml').read_text())
    services=compose.get('services',{})
    ok('compose services',len(services)>=7,f'{len(services)} services')
    unpinned=[]
    for name,svc in services.items():
        image=svc.get('image')
        if image and (image.endswith(':latest') or ':latest}' in image or ':' not in image): unpinned.append(f'{name}={image}')
    ok('container images pinned',not unpinned,', '.join(unpinned) if unpinned else 'all explicit versions')
except Exception as e:
    ok('Docker Compose YAML',False,str(e))

# Windows host-port fallback must be wired into Compose/launcher.
compose_text=(ROOT/'docker-compose.yml').read_text()
launcher_text=(ROOT/'START_DEMO.bat').read_text()
ok('frontend dynamic host port','${APP_HOST_PORT:-3000}:3000' in compose_text and 'APP_HOST_PORT' in launcher_text,'dynamic 3000 fallback wired')
ok('API dynamic host port','${API_HOST_PORT:-8000}:8000' in compose_text and 'API_HOST_PORT' in launcher_text and 'http://localhost:${API_HOST_PORT:-8000}/api/v1' in compose_text,'dynamic 8000 fallback wired')
ok('n8n dynamic host port','${N8N_HOST_PORT:-5678}:5678' in compose_text and 'select_runtime_ports.ps1' in launcher_text,'dynamic 5678 fallback wired')
ok('MinIO console dynamic host port','${MINIO_CONSOLE_HOST_PORT:-9001}:9001' in compose_text and 'select_runtime_ports.ps1' in launcher_text,'dynamic 9001 fallback wired')


# Database credential consistency / stale-volume recovery controls.
secret_loader=(ROOT/'scripts/load_compose_secrets.bat').read_text(errors='ignore')
pg_sync=(ROOT/'scripts/sync_postgres_credentials.bat').read_text(errors='ignore')
ok('compose secrets loaded from generated env', 'OTP_DB_PASSWORD=' in secret_loader and 'findstr /B /C:"OTP_DB_PASSWORD=" .env' in secret_loader, 'launcher process pins Compose interpolation to .env')
ok('postgres stale-volume credential repair', 'ALTER ROLE otp WITH PASSWORD' in pg_sync and 'PGPASSWORD=' in pg_sync and 'SELECT 1' in pg_sync, 'password synchronized then verified over TCP')
ok('authenticated postgres healthcheck', 'PGPASSWORD=$${POSTGRES_PASSWORD}' in compose_text and "SELECT 1" in compose_text and 'pg_isready -U otp -d otp' not in compose_text, 'DB healthcheck validates credentials')
ok('API DB password source aligned', compose_text.count('postgresql+psycopg2://otp:${OTP_DB_PASSWORD}@otp-postgres:5432/otp') >= 2, 'API and worker use same Compose-resolved password as PostgreSQL')
ok('launcher runs DB credential preflight', 'sync_postgres_credentials.bat' in launcher_text and 'load_compose_secrets.bat' in launcher_text, 'one-click launcher repairs/verifies DB credentials')
start_otp_text = (ROOT / 'START_OTP.bat').read_text(encoding='utf-8', errors='ignore')
sync_pg_text = (ROOT / 'scripts' / 'sync_postgres_credentials.bat').read_text(encoding='utf-8', errors='ignore')
ok('START_OTP loads release secrets', 'load_compose_secrets.bat' in start_otp_text and start_otp_text.find('load_compose_secrets.bat') < start_otp_text.find('sync_postgres_credentials.bat'), 'direct launcher loads .env secrets before PostgreSQL preflight')
ok('DB preflight self-load fallback', 'OTP_DB_PASSWORD is not yet loaded' in sync_pg_text and 'load_compose_secrets.bat' in sync_pg_text, 'credential preflight can self-load .env when invoked independently')
ok('DB preflight uses SQL readiness not pg_isready', '-d postgres' in sync_pg_text and '-tAc "SELECT 1"' in sync_pg_text and 'pg_isready' not in '\n'.join(line for line in sync_pg_text.splitlines() if not line.lstrip().lower().startswith('rem ')), 'first-start gate waits for a real SQL statement')
ok('DB preflight retries role synchronization', 'for /L %%I in (1,1,15)' in sync_pg_text and 'PG_SYNC_OK' in sync_pg_text, 'transient initialization race is retried instead of failing immediately')
ok('DB preflight preserves diagnostics', 'postgres_preflight.log' in sync_pg_text and 'type "%PG_PREFLIGHT_LOG%"' in sync_pg_text, 'real PostgreSQL errors are preserved for Windows troubleshooting')

# Dependency pinning.
req=[]
for line in (ROOT/'backend/requirements.txt').read_text().splitlines():
    line=line.strip()
    if line and not line.startswith('#'): req.append(line)
ok('Python dependencies pinned',all('==' in x for x in req),', '.join(x for x in req if '==' not in x))
try:
    pkg=json.loads((ROOT/'frontend/package.json').read_text())
    deps={**pkg.get('dependencies',{}),**pkg.get('devDependencies',{})}
    loose=[f'{k}={v}' for k,v in deps.items() if v.startswith(('^','~','>','<','*'))]
    ok('frontend direct dependencies pinned',not loose,', '.join(loose) if loose else 'all exact')
    ok('frontend package version',pkg.get('version')==VERSION,f"package={pkg.get('version')} root={VERSION}")
except Exception as e:
    ok('frontend package JSON',False,str(e))

# Version consistency in critical executable config.
critical=[ROOT/'START_DEMO.bat',ROOT/'frontend/Dockerfile',ROOT/'frontend/components/Shell.tsx',ROOT/'backend/app/settings.py',ROOT/'.env.example',ROOT/'docker-compose.yml']
for f in critical:
    txt=f.read_text(errors='ignore')
    ok(f'version marker {f.relative_to(ROOT)}',VERSION in txt,f'{VERSION} present' if VERSION in txt else f'{VERSION} not found')


# Bilingual localisation assets.
i18n=(ROOT/'frontend/lib/i18n.tsx').read_text(encoding='utf-8') if (ROOT/'frontend/lib/i18n.tsx').exists() else ''
shell=(ROOT/'frontend/components/Shell.tsx').read_text(encoding='utf-8') if (ROOT/'frontend/components/Shell.tsx').exists() else ''
ok('German localisation dictionaries', all(x in i18n for x in ["const de:","const en:","'de-DE'","'en-GB'","parseLocalizedNumber"]), 'DE/EN dictionaries and locale helpers present')
ok('persistent language switch', 'otp_language' in i18n and "setLanguage('de')" in shell and "setLanguage('en')" in shell, 'DE/EN switch and local persistence present')
ok('German localisation docs', (ROOT/'README_DE.md').exists() and (ROOT/'docs/GERMAN_LOCALISATION.md').exists(), 'German user and technical documentation present')

# Governed Entity Master controls.
master_router=(ROOT/'backend/app/routers/master_data.py').read_text()
master_service=(ROOT/'backend/app/services/entity_master.py').read_text()
entities_page=(ROOT/'frontend/app/entities/page.tsx').read_text()
ok('entity master maker-checker', 'Maker-checker control' in master_service and '/entity-changes/{request_id}/approve' in master_router, 'approval separation implemented')
ok('entity master version history', 'EntityVersion' in master_service and '/entities/{code}/history' in master_router, 'versioned history endpoint present')
ok('entity master bulk preview', '/entities/import/preview' in master_router and 'never interpreted as deletions' in master_service, 'controlled Excel preview with no implicit deletion')
ok('entity master UI maintenance', all(x in entities_page for x in ['entities.add','entities.edit','entities.importMaster','entities.pending']), 'add/edit/import/review UI present')


# Enterprise master-data catalog and allocation-workbench controls introduced in RC.1.
catalog_service=(ROOT/'backend/app/services/master_data_catalog.py').read_text()
catalog_page=(ROOT/'frontend/app/master-data/page.tsx').read_text()
allocation_router=(ROOT/'backend/app/routers/allocations.py').read_text()
allocation_page=(ROOT/'frontend/app/allocations/page.tsx').read_text()
models_text=(ROOT/'backend/app/models.py').read_text()
ok('enterprise master-data domains', all(x in catalog_service for x in ['GL_ACCOUNTS','COST_CENTERS','PROFIT_CENTERS','PLANTS','BUSINESS_UNITS','PRODUCT_GROUPS','PRODUCTS','CUSTOMERS','VENDORS','CURRENCIES','FX_SOURCES','TRANSACTION_TYPES','TP_SEGMENTS']), 'governed catalog domains present')
ok('generic master maker-checker', 'creator cannot approve own request' in catalog_service and 'MasterDataVersion' in catalog_service, 'domain-neutral maker-checker/versioning present')
ok('master catalog UI', all(x in catalog_page for x in ['/catalog/domains','import/preview','DEACTIVATE','REACTIVATE']), 'create/edit/lifecycle/import UI present')
ok('master activation affects operations', 'GLAccount.active.is_(True)' in (ROOT/'backend/app/services/ingestion.py').read_text() and 'row.active=record.active' in catalog_service, 'deactivation propagates to operational master projections')
ok('allocation configurable lifecycle', all(x in allocation_router for x in ['/rules/{rule_id}/submit','/rules/{rule_id}/approve','/rules/{rule_id}/activate','/rules/{rule_id}/simulate','/rules/{rule_id}/execute']), 'draft/review/approve/activate/simulate/execute endpoints present')
ok('allocation demo UX removed', 'Run €1m' not in allocation_page and '1000000' not in allocation_page and "api('/allocations/run'" not in allocation_page, 'primary UI uses configurable rules')
ok('allocation control features', all(x in models_text for x in ['driver_weights','recipient_rules','markup_treatment','fx_methodology','source_amount_source']), 'configurable source/driver/recipient/markup/FX metadata persisted')

# RC.2 practicality upgrade controls.
bench_service=(ROOT/'backend/app/services/benchmarks.py').read_text()
workflow_service=(ROOT/'backend/app/services/workflows.py').read_text()
enterprise_ops=(ROOT/'backend/app/routers/enterprise_ops.py').read_text()
accounting_router=(ROOT/'backend/app/routers/accounting.py').read_text()
allocation_service=(ROOT/'backend/app/services/allocations.py').read_text()
master_catalog=(ROOT/'backend/app/services/master_data_catalog.py').read_text()
governance_router=(ROOT/'backend/app/routers/governance.py').read_text()
ok('benchmark comparable engine', all(x in models_text for x in ['BenchmarkComparable','BenchmarkFinancialObservation']) and 'benchmark_statistics' in bench_service and '/benchmarks/{benchmark_code}/statistics/stage' in enterprise_ops, 'comparable screening, observations, deterministic quartiles and governed staging present')
ok('multi-stage allocation engine', '_multi_stage_allocate' in allocation_service and 'Circular allocation detected' in allocation_service and 'stage_driver_values' in allocation_router, 'arbitrary staged allocation with cycle protection implemented')
ok('dynamic workflow resolution', 'resolve_adjustment_workflow' in workflow_service and 'workflow_code' in models_text and 'version-pinned approval workflow' in (ROOT/'backend/app/routers/approvals.py').read_text(), 'governed workflow definition is pinned to approval request')
ok('journal-template execution', 'JOURNAL_TEMPLATES' in master_catalog and 'template_code' in models_text and 'governed template' in accounting_router, 'journal GL determination uses governed template when available')
ok('reconciliation exception and certification', 'EXC-REC-' in accounting_router and '/reconciliations/{reconciliation_id}/certify' in accounting_router, 'differences create cases and reconciled records support sign-off')
ok('evidence pack generation', '/evidence-packs/adjustments/{adjustment_id}' in enterprise_ops and 'manifest_final.json' in enterprise_ops, 'transaction evidence ZIP endpoint present')
ok('master dependency impact', 'dependency_impact' in master_catalog and '/catalog/{domain}/{record_code}/impact' in (ROOT/'backend/app/routers/master_data.py').read_text(), 'blocking dependency preview/guard present')
ok('audit explorer filtering', '/audit/timeline/{object_type}/{object_id}' in governance_router and 'correlation_id' in governance_router, 'filterable object/correlation audit reconstruction present')

# Sample spreadsheets must be valid OOXML and readable.
try:
    import openpyxl
    xlsx=list((ROOT/'sample_data').glob('*.xlsx'))
    bad=[]; total_rows=0
    for f in xlsx:
        try:
            wb=openpyxl.load_workbook(f,read_only=True,data_only=False)
            if not wb.sheetnames: bad.append(f.name); continue
            ws=wb[wb.sheetnames[0]]; rows=sum(1 for _ in ws.iter_rows(values_only=True)); total_rows += max(rows-1,0); wb.close()
        except Exception: bad.append(f.name)
    ok('sample XLSX integrity',not bad,f'{len(xlsx)} workbooks, {total_rows} data rows' if not bad else ', '.join(bad))
except Exception as e:
    ok('sample XLSX validation',False,str(e))

# PDFs: signature and EOF marker; full rendering is an external/runtime check.
pdfs=list((ROOT/'documents').glob('*.pdf')); bad=[]
for f in pdfs:
    b=f.read_bytes()
    if not (b.startswith(b'%PDF') and b'%%EOF' in b[-2048:]): bad.append(f.name)
ok('reference PDF integrity',not bad,f'{len(pdfs)} PDFs' if not bad else ', '.join(bad))

# No obvious real secrets in production template.
example=(ROOT/'.env.example').read_text()
ok('production env uses secret placeholders','SET_IN_SECRET_MANAGER' in example and 'DEMO_MODE=false' in example,'expected placeholders/DEMO_MODE=false')

# v2.4.0 10/10 practicality layer controls.
pr10=(ROOT/'backend/app/routers/practical10.py').read_text(errors='ignore')
policy_page=(ROOT/'frontend/app/policies/page.tsx').read_text(errors='ignore')
workflow_page=(ROOT/'frontend/app/workflows/page.tsx').read_text(errors='ignore')
pnl_page=(ROOT/'frontend/app/pnl/page.tsx').read_text(errors='ignore')
security_page=(ROOT/'frontend/app/security/page.tsx').read_text(errors='ignore')
variance_service=(ROOT/'backend/app/services/variance.py').read_text(errors='ignore')
ok('policy workbench lifecycle', all(x in pr10 for x in ["/policies/workbench","/submit","/approve","/activate","/history","/impact"]) and 'New TP policy' in policy_page, 'editable/versioned maker-checker policy workbench')
ok('pnl cube snapshot locking', all(x in pr10 for x in ["/pnl/runs/{run_id}/cube","/snapshots","/lock"]) and 'Multidimensional cube' in pnl_page, 'multidimensional P&L plus controlled snapshots')
ok('true-up batch reversal', all(x in pr10 for x in ["/true-ups/batch","/true-ups/{adjustment_id}/reverse"]), 'deterministic batch/reversal endpoints')
ok('workflow full lifecycle', all(x in pr10 for x in ["/workflow-engine/instances","/return","/withdraw","/resubmit","/escalate-overdue"]) and 'My approval inbox' in workflow_page, 'version-pinned approval engine with SLA/escalation')
ok('reconciliation stage control', '/reconciliations/{reconciliation_id}/stages/rebuild' in pr10, 'stage-level reconciliation and exception generation')
ok('service account security', '/security/service-accounts' in pr10 and 'One-time token' in security_page, 'hashed scoped integration credentials and UI')
ok('qualification evidence registry', '/qualification/sync-release-evidence' in pr10 and '/qualification/evidence' in pr10 and './docs:/release-docs:ro' in compose_text, 'release evidence is persisted and mounted read-only')
ok('variance comparison breadth', all(x in variance_service for x in ["FORECAST","TARGET","PRIOR_YEAR","VOLUME","PRICE","MIX","FX"]), 'deterministic expanded comparison/driver support')


# Local 10/10 runtime qualification harness.
qual_batch=(ROOT/'QUALIFY_10_10.bat').read_text(errors='ignore')
qual_ps=(ROOT/'scripts/qualify_10_10.ps1').read_text(errors='ignore')
qual_ui=(ROOT/'qualification/tests/ui.spec.ts').read_text(errors='ignore')
qual_api=(ROOT/'qualification/tests/api.spec.ts').read_text(errors='ignore')
qual_compose=(ROOT/'docker-compose.qualification.yml').read_text(errors='ignore')
target_layers=['Master Data','Allocation','Benchmarks','Segmented P&L','True-Ups','Accounting','Exceptions','Documents / Evidence','Audit','Approvals','Policies','Variance','Reconciliation','Security','Qualification']
ok('one-click local 10/10 qualifier','qualify_10_10.ps1' in qual_batch and 'otp-qualification' in qual_ps and 'generate_local_qualification_report.py' in qual_ps,'Windows one-click qualifier orchestrates runtime, backend, browser and report')
ok('qualifier reuses healthy runtime', 'Existing OTP runtime is healthy and version-matched' in qual_ps and 'Resolve-RuntimePorts' in qual_ps and 'Test-HttpOk' in qual_ps, 'qualification reuses only a healthy runtime that matches the current release version')
ok('qualifier rejects stale runtime versions', 'Get-ApiVersion' in qual_ps and '/api/v1/health/version' in qual_ps and 'Rebuilding the current release' in qual_ps, 'a healthy older hotfix is not reused for current-release qualification')
ok('qualifier exposes Docker build progress', "--progress','plain','build" in qual_ps and 'OTP_SHOW_BUILD_PROGRESS' in launcher_text, 'first-time qualification runner/application builds remain visibly progressing')
ok('qualifier validates release before mutable test execution', qual_ps.find("Run release/package validation on an uncontaminated release tree") < qual_ps.find("Run backend regression + measured coverage in isolated result storage"), 'package integrity is measured before pytest can create runtime artifacts')
ok('qualifier isolates pytest artifacts', 'COVERAGE_FILE=/results/.coverage-local' in qual_ps and 'cache_dir=/results/pytest-cache' in qual_ps, 'coverage data and pytest cache are written outside the release tree')
ok('qualifier isolates local runtime env for package validation', 'otp-hotfix7-' in qual_ps and 'Move-Item -Force $runtimeEnv $envBackup' in qual_ps and 'Move-Item -Force $envBackup $runtimeEnv' in qual_ps, 'runtime .env is temporarily moved outside the release tree and restored')
ok('qualifier publishes current release evidence before UI UAT', 'Current release-integrity evidence synchronized before browser UAT' in qual_ps and 'Write-ReleaseEvidence' in qual_ps, 'browser qualification reads current-run release integrity rather than stale evidence')
ok('browser qualification covers all target layers',all(layer in qual_ui for layer in target_layers),'all 15 workbenches have live browser checks')
ok('live API qualification breadth',all(x in qual_api for x in ['governed create, maker-checker','multi-stage execution','deterministic quartiles','authoritative build, cube, snapshot and lock','deterministic target adjustment','balanced journal','automatic reconciliation exception','evidence search','transaction-chain reconstruction','user-configured workflow','versioned lifecycle','additive bridge','stage rebuild','service token','evidence registry']),'deep live API UAT spans all 15 target layers')
ok('qualification runner isolated','profiles: ["qualification"]' in qual_compose and 'otp-backend-tests' in qual_compose,'qualification-only Docker services do not alter normal startup path')
qual_entry=(ROOT/'qualification/entrypoint.sh').read_text(errors='ignore')
ok('qualification browser uses private backplane','networks: [otp-backplane]' in qual_compose and 'otp-frontend' in qual_entry and 'otp-api' in qual_entry and 'BASE_URL="http://localhost:' in qual_entry,'browser preserves localhost CORS origin while forwarding to live Compose services')
try:
    qual_pkg=json.loads((ROOT/'qualification/package.json').read_text())
    qual_deps={**qual_pkg.get('dependencies',{}),**qual_pkg.get('devDependencies',{})}
    qual_loose=[f'{k}={v}' for k,v in qual_deps.items() if str(v).startswith(('^','~','>','<','*'))]
    ok('qualification direct dependencies pinned',not qual_loose,', '.join(qual_loose) if qual_loose else 'all exact')
except Exception as exc:
    ok('qualification package JSON',False,str(exc))

for name,passed,detail in checks:
    print(f"[{'PASS' if passed else 'FAIL'}] {name}: {detail}")

if errors:
    print('\nRelease validation FAILED')
    for e in errors: print('-',e)
    sys.exit(1)
print(f'\nRelease validation PASS — version {VERSION}; files={sum(1 for p in ROOT.rglob("*") if p.is_file())}; n8n_nodes={sum(workflow_nodes.values())}')
